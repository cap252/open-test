"""Portable research evidence, kept separate from locally attested run records."""

import csv
import json
import math
import re
from collections import Counter
from pathlib import Path
from ._io import read_json, atomic_json, sha256
from .results import table_text
from .configuration import load_config, validate
from .models.spec import model_names
from .models.evaluation_profiles import MODELS, HISTORICAL

IDS = dict(
    zip(
        "01 02 03 04 05 06 07 08 09 10 11 12 13 14 15 16 17 18 19 20 22 27 29 37".split(),
        "resnet18 resnet50 resnet101 mobilenet_v1 mobilenet_v2 mobilenet_v3_large efficientnet_b0 efficientnet_lite4 squeezenet1_1 shufflenet_v2_x1_0 vgg16 inception_v3 yolov3_tiny yolov4_tiny yolov5n yolov5s yolov7_tiny yolov8n ssd_mobilenet_v2_320 nanodet_plus_m_320 retinanet_resnet50_fpn retinaface_mnet025_640 posenet_101 deeplabv3_mnv2_voc".split(),
    )
)
IDS["efficientnet_lite0_timm_in1k"] = "efficientnet_lite0"
FAMILIES = {"asym_pc_float": "float", "asym_pc_pot": "pot", "sym_pt_pot": "sym_pot"}


class Portable:
    def __init__(self):
        self.bindings = {}

    def path(self, value):
        if value in self.bindings:
            return self.bindings[value]
        for marker, name in (
            ("/midap_quantization/runs/", "research-runs"),
            ("/midap_cnn_quant/", "research-repo"),
            ("/onnx-quantization/result/", "research-snapshots"),
        ):
            if marker in value:
                result = "artifact://" + name + "/" + value.split(marker, 1)[1]
                break
        else:
            import hashlib

            result = (
                "artifact://external/"
                + hashlib.sha256(value.encode()).hexdigest()[:16]
                + "/"
                + Path(value).name
            )
        self.bindings[value] = result
        return result

    def __call__(self, value):
        if isinstance(value, dict):
            return {self(k): self(v) for k, v in value.items()}
        if isinstance(value, list):
            return [self(v) for v in value]
        if isinstance(value, str):
            if value.startswith(("{", "[")):
                try:
                    return json.dumps(
                        self(json.loads(value)), ensure_ascii=False, sort_keys=True
                    )
                except json.JSONDecodeError:
                    pass
            return re.sub(
                r"/(?:mnt|home|tmp)/[^\s\"'<>;,\]\)]+", lambda m: self.path(m[0]), value
            )
        return value


def validate_selection(rows):
    keys = [(r["model_id"], r.get("family") or "", r["method"]) for r in rows]
    if (
        len(rows) != 273
        or len(set(keys)) != 273
        or len({r["model_id"] for r in rows}) != 39
    ):
        raise ValueError("Expected unique 39 x 7 selected records")
    means = []
    for row in rows:
        if row.get("snapshot_state", row["state"]) not in (
            "completed",
            "evaluated",
            "reused",
        ):
            continue
        fp = row["fp32_accuracy"]
        score = row.get(
            "accuracy", fp if row["method"] == "fp32" else row.get("int8_accuracy")
        )
        if not all(
            isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)
            for v in (fp, score)
        ):
            raise ValueError("Missing/nonfinite completed metric")
        unit = row.get("unit", row.get("display_unit"))
        if unit == "dB":
            if row.get("recovery_percent") is not None:
                raise ValueError("PSNR must not have percentage recovery")
            if row["method"] != "fp32" and not math.isclose(
                row["delta_db"], score - fp, abs_tol=1e-9
            ):
                raise ValueError("Invalid dB delta")
        elif unit == "%":
            if fp <= 0 or not math.isclose(
                row["recovery_percent"], 100 * score / fp, abs_tol=1e-8
            ):
                raise ValueError("Invalid measured-FP32 recovery")
            if row["method"] != "fp32" and not math.isclose(
                row["drop_pp"], fp - score, abs_tol=1e-8
            ):
                raise ValueError("Invalid percentage-point drop")
        else:
            raise ValueError("Unknown research metric unit")
    for family in FAMILIES:
        group = [
            r
            for r in rows
            if r.get("family") == family and r.get("unit", r.get("display_unit")) == "%"
        ]
        for method in ("rtn", "adaround", "cross"):
            by_model = {}
            for row in group:
                if method != "cross" and row["method"] != method:
                    continue
                if row.get("snapshot_state", row["state"]) in (
                    "completed",
                    "evaluated",
                    "reused",
                ):
                    mid = row["model_id"]
                    score = row["recovery_percent"]
                    by_model[mid] = max(score, by_model.get(mid, -math.inf))
            means.append(
                dict(
                    family=family,
                    method=method,
                    completed_models=len(by_model),
                    expected_models=38,
                    mean_recovery_percent=sum(by_model.values()) / 38
                    if len(by_model) == 38
                    else None,
                    observed_partial_mean_recovery_percent=sum(by_model.values())
                    / len(by_model)
                    if by_model and len(by_model) < 38
                    else None,
                    selection="posthoc RTN/AdaRound maximum"
                    if method == "cross"
                    else "frozen current selection",
                )
            )
    return means


def export_tables(tables, directory, *, xlsx=False):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    counts = {}
    for name, rows in tables.items():
        fields = list(dict.fromkeys(k for r in rows for k in r))
        for suffix, delimiter in (("csv", ","), ("tsv", "\t")):
            path = directory / (name + "." + suffix)
            path.write_text(
                table_text(rows, fields, delimiter=delimiter), encoding="utf-8"
            )
            parsed = list(csv.DictReader(path.open(), delimiter=delimiter))
            if len(parsed) != len(rows):
                raise ValueError("Export row count changed")
            for original, back in zip(rows, parsed):
                for key, value in original.items():
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        if float(back[key]) != value:
                            raise ValueError("Export numeric value changed")
        counts[name] = len(rows)
    atomic_json(
        directory / "records.json", dict(origin="external_research", tables=tables)
    )
    if xlsx:
        from openpyxl import Workbook, load_workbook

        workbook = Workbook()
        workbook.remove(workbook.active)
        for name, rows in tables.items():
            sheet = workbook.create_sheet(name[:31])
            fields = list(dict.fromkeys(k for r in rows for k in r))
            sheet.append(fields)
            for row in rows:
                sheet.append(
                    [
                        json.dumps(row.get(k), ensure_ascii=False)
                        if isinstance(row.get(k), (dict, list))
                        else row.get(k)
                        for k in fields
                    ]
                )
        workbook.save(directory / "results.xlsx")
        actual = load_workbook(
            directory / "results.xlsx", read_only=True, data_only=True
        )
        for name, rows in tables.items():
            values = list(actual[name[:31]].values)
            fields = values[0]
            for row, cells in zip(rows, values[1:]):
                for key, value in zip(fields, cells):
                    original = row.get(key)
                    if (
                        isinstance(original, (int, float))
                        and not isinstance(original, bool)
                        and not math.isclose(
                            value, original, rel_tol=1e-14, abs_tol=1e-12
                        )
                    ):
                        raise ValueError("XLSX numeric value changed")
        actual.close()
    return counts


def selected_config(rows):
    cfg = load_config("cnn38")
    cfg["name"] = "research_selected_20261017"
    cfg["conditions"] = {"fp32": {}}
    old = cfg["schemes"]
    cfg["schemes"] = {}
    cfg["percentiles"] = {m: {} for m in cfg["models"]}
    evidence = []
    for family, key in FAMILIES.items():
        for method in ("rtn", "adaround"):
            name = family + "_" + method
            cfg["schemes"][name] = dict(old[key])
            cfg["conditions"][name] = dict(
                scheme=name, **({"adaround": True} if method == "adaround" else {})
            )
            for row in rows:
                mid = IDS.get(row["model_id"], row["model_id"])
                if (
                    mid not in cfg["models"]
                    or row.get("family") != family
                    or row["method"] != method
                ):
                    continue
                cfg["percentiles"][mid][name] = row["percentile"]
                evidence.append(
                    dict(
                        model=mid,
                        condition=name,
                        source_condition=row["condition"],
                        percentile=row["percentile"],
                        evidence_sha256=row["evidence_sha256"],
                        selection_basis=row.get("selection_basis")
                        or row.get("report_origin"),
                    )
                )
    cfg["evaluation"]["profiles"] = {m: HISTORICAL for m in sorted(MODELS)}
    # Training settings remain explicit; this config does not attest cross-environment retraining.
    validate(cfg)
    return cfg, evidence


def recorded_training(rows, sources):
    """Verify lightweight original results/audits without claiming payload replay."""
    roots = set()
    for source in sources:
        value = source["source"]
        for marker in ("/runs/", "/midap_cnn_quant/results/"):
            if marker in value:
                roots.add(Path(value.split(marker, 1)[0] + marker))
    records = []
    for row in rows:
        if row["method"] != "adaround":
            continue
        evidence = Path(row["evidence"])
        candidates = (
            [evidence]
            if evidence.is_absolute()
            else [root / str(row["source_run"]) / evidence for root in sorted(roots)]
        )
        path = next((p for p in candidates if p.is_file()), None)
        item = dict(
            model=IDS.get(row["model_id"], row["model_id"]),
            condition=row["condition"],
            percentile=row["percentile"],
        )
        if path is None:
            records.append(
                dict(
                    item, state="missing_result", recorded_sha256=row["evidence_sha256"]
                )
            )
            continue
        if sha256(path) != row["evidence_sha256"]:
            raise ValueError("Selected measurement changed: " + str(path))
        measured = read_json(path)
        audit_path = Path(measured.get("audit_path") or path.parent / "audit.json")
        if not audit_path.is_file():
            records.append(
                dict(item, state="missing_audit", result_sha256=sha256(path))
            )
            continue
        if sha256(audit_path) != measured["audit_sha256"]:
            raise ValueError("Selected AdaRound audit changed")
        audit = read_json(audit_path)
        ar = audit["adaround"]
        records.append(
            dict(
                item,
                state="verified_result_and_audit",
                result_artifact=str(path),
                result_sha256=sha256(path),
                audit_artifact=str(audit_path),
                audit_sha256=sha256(audit_path),
                config=ar["config"],
                sample_count=ar["sample_count"],
                samples_identity=ar["samples_identity"],
                frozen_encoding_identity=ar["frozen_encoding_identity"],
                parameter_identity=audit["parameter_identity"],
                model_sha256_recorded=measured["model_sha256"],
                model_payload_reverified=False,
                scheme=measured.get("scheme"),
                evaluation_samples=measured["result"]["sample_count"],
            )
        )
    return records


def training_configs(config, records):
    """Use existing per-run settings; do not add a second training execution path."""
    import copy

    configs = {}
    for model in config["models"]:
        matches = [r for r in records if r["model"] == model]
        if len(matches) != 3 or any(
            r["state"] != "verified_result_and_audit" for r in matches
        ):
            continue
        if any(r["config"] != matches[0]["config"] for r in matches):
            raise ValueError("Per-condition training configurations differ: " + model)
        cfg = copy.deepcopy(config)
        cfg["models"] = [model]
        cfg["name"] += "_" + model
        cfg["percentiles"] = {model: cfg["percentiles"][model]}
        cfg["evaluation"]["profiles"] = {
            k: v for k, v in cfg["evaluation"]["profiles"].items() if k == model
        }
        cfg["adaround"] = matches[0]["config"]
        cfg["runtime"]["device"] = "cuda:0"
        validate(cfg)
        configs[model] = cfg
    return configs


def _import_snapshot(source, output, *, bindings_output, xlsx=False):
    source, output = Path(source), Path(output)
    if output.exists():
        raise FileExistsError(output)
    data = read_json(source / "all_records.json")
    tables = data["tables"]
    rows = tables["All_Selected_Current"]
    means = validate_selection(rows)
    recorded = {(r["family"], r["method"]): r for r in tables["Overall_Current"]}
    for row in means:
        key = (
            row["family"],
            "cross_best" if row["method"] == "cross" else row["method"],
        )
        if key not in recorded:
            matches = [
                r
                for r in tables["Overall_Current"]
                if r["family"] == row["family"]
                and r["method"] not in ("rtn", "adaround")
            ]
            prior = (
                matches[0] if row["method"] == "cross" and len(matches) == 1 else None
            )
        else:
            prior = recorded[key]
        if (
            not prior
            or prior["mean_recovery_percent"] != row["mean_recovery_percent"]
            and not math.isclose(
                prior["mean_recovery_percent"],
                row["mean_recovery_percent"],
                abs_tol=1e-10,
            )
        ):
            raise ValueError("Recomputed cohort mean differs")
    gate = read_json(source / "rtn_completion_verification.json")
    rtn = tables["RTN_Followup_Raw"]
    if not gate["ready"] or gate["completed"] != 90 or len(rtn) != 90:
        raise ValueError("RTN90 is incomplete")
    portable = Portable()
    proof = []
    verified_sources = []
    for row in tables["All_Sources"]:
        path = source / row["snapshot"]
        if not path.resolve().is_relative_to(source.resolve()):
            raise ValueError("Unsafe source snapshot path")
        stored_hash = sha256(path)
        content_hash = stored_hash
        if row.get("compression") == "gzip":
            import gzip, hashlib

            content_hash = hashlib.sha256(
                gzip.decompress(path.read_bytes())
            ).hexdigest()
        if content_hash != row["sha256"]:
            raise ValueError("Captured source changed: " + str(row["snapshot"]))
        verified_sources.append(dict(row, stored_sha256=stored_hash))
    for row in rtn:
        path = Path(row["evidence"])
        if sha256(path) != row["evidence_sha256"]:
            raise ValueError("RTN evidence changed")
        measured = read_json(path)
        checks = {}
        for kind in ("model", "audit"):
            if kind == "audit" and row["condition"] == "fp32":
                continue
            asset = Path(measured[kind + "_path"])
            if sha256(asset) != measured[kind + "_sha256"]:
                raise ValueError("RTN payload changed: " + kind)
            checks[kind] = dict(
                artifact=portable.path(str(asset)),
                sha256=measured[kind + "_sha256"],
                reverified=True,
            )
        reused = measured.get("reuse")
        if reused:
            if sha256(reused["source_path"]) != reused["source_sha256"]:
                raise ValueError("RTN reuse source changed")
            checks["reuse"] = dict(
                artifact=portable.path(reused["source_path"]),
                sha256=reused["source_sha256"],
                reverified=True,
            )
        if measured["result"]["sample_count"] != row["samples"]:
            raise ValueError("RTN sample count differs")
        proof.append(
            dict(
                model=row["model_id"],
                condition=row["condition"],
                state=row["state"],
                checks=checks,
            )
        )
    supported = set(model_names())
    selected = []
    for row in rows:
        mid = IDS.get(row["model_id"], row["model_id"])
        value = dict(
            row,
            runtime_model=mid,
            runtime_supported=mid in supported,
            locally_reproduced=False,
            evaluation_profile=HISTORICAL
            if mid in MODELS
            else "recorded_source_profile",
            verification_level="recorded measurement and snapshot hashes; full model payload not reverified",
        )
        selected.append(portable(value))
    config, selection = selected_config(rows)
    training = recorded_training(rows, tables["All_Sources"])
    configs = training_configs(config, training)
    output.mkdir(parents=True)
    selected_tables = dict(
        Selected=selected,
        RTN_Followup=portable(rtn),
        Overall=means,
        Sources=portable(verified_sources),
        Coverage=portable(tables["Coverage_All"]),
        YOLO_Size_Comparison=portable(tables["YOLO_Size_Comparison"]),
        Blocked=portable(tables["ArcFace_Blocked"]),
    )
    counts = export_tables(selected_tables, output, xlsx=xlsx)
    execution = []
    for path in sorted((source / "execution_sources").glob("*.json")):
        doc = read_json(path)
        execution.append(
            dict(artifact="execution_sources/" + path.name, source_sha256=sha256(path))
        )
        if "source_files" in doc:
            root = Path(doc["source"])
            for relative, digest in doc["source_files"].items():
                if sha256(root / relative) != digest:
                    raise ValueError("Measured worker source changed: " + relative)
            diff = source / "execution_sources" / (path.stem + ".diff")
            if sha256(diff) != doc["diff_sha256"]:
                raise ValueError("Worker diff changed")
        elif path.name == "runner_binding.json":
            if sha256(doc["source"]) != doc["sha256"]:
                raise ValueError("Actual runner changed")
        atomic_json(output / "execution_sources" / path.name, portable(doc))
    atomic_json(
        output / "selected_policy.json",
        dict(
            percentiles=config["percentiles"],
            schemes=config["schemes"],
            conditions=config["conditions"],
            evaluation=config["evaluation"],
        ),
    )
    atomic_json(output / "recorded_training.json", portable(training))
    for model, cfg in configs.items():
        atomic_json(output / "configs" / (model + ".json"), cfg)
    atomic_json(
        output / "selection_manifest.json",
        dict(
            policy_sha256=sha256(output / "selected_policy.json"),
            configs={
                model: dict(
                    path="configs/" + model + ".json",
                    sha256=sha256(output / "configs" / (model + ".json")),
                )
                for model in configs
            },
            source_sha256=sha256(source / "all_records.json"),
            selections=portable(selection),
            policy="Current selected table only. RTN followup maxima do not replace historical AdaRound choices.",
            retraining_equivalence="unverified; per-model recorded settings and frozen assets remain authoritative",
        ),
    )
    states = Counter(r["state"] for r in rtn)
    verification = dict(
        origin="external_research",
        captured_at=data["captured_at"],
        source_sha256=sha256(source / "all_records.json"),
        model_count=39,
        runtime_supported=38,
        selected_slots=273,
        rtn90=dict(states=states, payloads=proof),
        source_snapshot_hashes_verified=len(tables["All_Sources"]),
        selected_adaround_result_audits_verified=sum(
            r["state"] == "verified_result_and_audit" for r in training
        ),
        runnable_model_configs=len(configs),
        execution_sources=execution,
        table_counts=counts,
        means_recomputed=True,
        local_full_accuracy=False,
        stages=dict(
            results_import="passed",
            same_output_postprocess="see local validation evidence",
            frozen_execution="see local validation evidence",
            new_environment_retraining="not_run",
            full_accuracy="not_run",
            rtx200="pending",
        ),
    )
    atomic_json(output / "verification.json", verification)
    atomic_json(
        bindings_output,
        dict(artifact_roots=portable.bindings, source_snapshot=str(source.resolve())),
    )
    if re.search(r"/(?:mnt|home|tmp)/", (output / "records.json").read_text()):
        raise ValueError("Local absolute path leaked into public records")
    return verification


def import_snapshot(source, output, *, bindings_output, xlsx=False):
    import tempfile

    output = Path(output)
    if output.exists():
        raise FileExistsError(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        dir=output.parent, prefix=".research-import-"
    ) as temporary:
        staged = Path(temporary) / "snapshot"
        bindings = Path(temporary) / "bindings.json"
        result = _import_snapshot(source, staged, bindings_output=bindings, xlsx=xlsx)
        atomic_json(bindings_output, read_json(bindings))
        staged.rename(output)
    return result
